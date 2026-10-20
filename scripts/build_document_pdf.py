#!/usr/bin/env python3
"""Build an RZCN documentation identification PDF from DOCX or text."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
from pathlib import Path
import re

import pymupdf
from docx import Document


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def wrap_line(line: str, width: int) -> list[str]:
    if len(line) <= width:
        return [line]
    return [line[i : i + width] for i in range(0, len(line), width)]


def read_input(path: Path) -> list[str]:
    if path.suffix.lower() == ".docx":
        document = Document(path)
        values = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                values.extend(cell.text.strip() for cell in row.cells if cell.text.strip())
        return values
    return [line.strip() for line in path.read_text(encoding="utf-8-sig").splitlines() if line.strip()]


def clean_input(
    values: list[str],
    drop_toc: bool,
    replacements: dict[str, str],
) -> list[str]:
    cleaned = []
    for value in values:
        for old, new in replacements.items():
            value = value.replace(old, new)
        if drop_toc and value == "目录":
            continue
        if drop_toc and re.fullmatch(r"[一二三四五六七八九十]+、[^…]+…+\s*\d+", value):
            continue
        cleaned.append(value)
    return cleaned


def select_rows(rows: list[str], mode: str, rows_per_page: int) -> tuple[list[str], str, bool]:
    page_count = math.ceil(len(rows) / rows_per_page)
    if mode == "auto":
        mode = "full" if page_count < 60 else "front-back"
    if mode == "full":
        return rows, mode, len(rows) % rows_per_page != 0
    page_rows = 30 * rows_per_page
    if len(rows) < page_rows:
        raise RuntimeError(f"front-back mode needs at least {page_rows} rows")
    return rows[:page_rows] + rows[-page_rows:], mode, False


def draw_header(page: pymupdf.Page, page_no: int, total: int, software: str, version: str, title: str) -> None:
    page.insert_text((28, 18), f"{software} {title} {version}", fontname="china-s", fontsize=9)
    page.insert_text((445, 18), f"第 {page_no} 页 / 共 {total} 页", fontname="china-s", fontsize=9)
    page.draw_line((28, 23), (567, 23), color=(0.72, 0.72, 0.72), width=0.35)


def build_pdf(args: argparse.Namespace) -> dict:
    replacements = {}
    if args.replace_json:
        replacements = json.loads(args.replace_json.read_text(encoding="utf-8"))
    values = clean_input(read_input(args.input), args.drop_toc, replacements)
    rows = [part for value in values for part in wrap_line(value, args.wrap_width)]
    if args.appendix:
        appendix = [line.strip() for line in args.appendix.read_text(encoding="utf-8-sig").splitlines() if line.strip()]
        rows.extend(part for line in appendix for part in wrap_line(line, args.wrap_width))
    if not rows:
        raise RuntimeError("document input is empty")

    selected, mode, partial = select_rows(rows, args.mode, args.rows_per_page)
    if partial and not args.allow_partial_last_page:
        raise RuntimeError("document has a final page below the row minimum; add factual appendix content or pass --allow-partial-last-page after review")
    pages = math.ceil(len(selected) / args.rows_per_page)
    if pages > 60:
        raise RuntimeError(f"selected document material has {pages} pages, exceeds 60")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc = pymupdf.open()
    doc.set_metadata(
        {
            "title": f"{args.software} {args.title} {args.version}",
            "author": args.software,
            "subject": "计算机软件著作权登记文档鉴别材料",
            "creator": args.software,
            "producer": args.software,
        }
    )
    for page_index in range(pages):
        page = doc.new_page(width=595, height=842)
        draw_header(page, page_index + 1, pages, args.software, args.version, args.title)
        y = 45
        page_rows = selected[page_index * args.rows_per_page : (page_index + 1) * args.rows_per_page]
        for local_index, line in enumerate(page_rows, start=1):
            page.insert_text((28, y), f"{local_index:02d}", fontname="china-s", fontsize=7)
            page.insert_text((52, y), line, fontname="china-s", fontsize=args.font_size)
            y += args.row_height
        marker = "完整文档" if mode == "full" else ("连续前30页" if page_index < 30 else "连续后30页")
        page.insert_text((28, 830), f"{marker}    每页至少 {args.rows_per_page} 行", fontname="china-s", fontsize=8)
    doc.save(args.output, garbage=4, deflate=True)
    doc.close()

    result = {
        "output": str(args.output),
        "software": args.software,
        "version": args.version,
        "title": args.title,
        "mode": mode,
        "pages": pages,
        "rows_per_page": args.rows_per_page,
        "partial_last_page": partial,
        "input_paragraphs": len(values),
        "visual_rows": len(rows),
        "selected_rows": len(selected),
        "sha256": sha256(args.output),
    }
    if args.summary:
        args.summary.parent.mkdir(parents=True, exist_ok=True)
        args.summary.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--software", required=True)
    parser.add_argument("--version", default="V1.0")
    parser.add_argument("--title", default="软件说明书")
    parser.add_argument("--mode", choices=["auto", "full", "front-back"], default="auto")
    parser.add_argument("--rows-per-page", type=int, default=30)
    parser.add_argument("--wrap-width", type=int, default=56)
    parser.add_argument("--font-size", type=float, default=8.2)
    parser.add_argument("--row-height", type=float, default=22)
    parser.add_argument("--drop-toc", action="store_true")
    parser.add_argument("--replace-json", type=Path)
    parser.add_argument("--appendix", type=Path)
    parser.add_argument("--allow-partial-last-page", action="store_true")
    parser.add_argument("--summary", type=Path)
    args = parser.parse_args()
    print(json.dumps(build_pdf(args), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
